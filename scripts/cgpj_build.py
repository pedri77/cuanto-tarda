#!/usr/bin/env python3
"""Construye los agregados de data/ a partir de los crudos cacheados en raw/.

Fuentes:
  - CGPJ "Series estadísticas" (XLSX): asuntos ingresados/resueltos/en trámite y tasas
    de congestión, resolución, pendencia y litigiosidad, 2001-2025, nacional + CCAA
    (+ provincias para total jurisdiccional).
  - INE: población a 1 de enero (solo para el derivado "pendientes por habitante").

Ejecutar tras scripts/cgpj_fetch.py: python3 scripts/cgpj_build.py
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW, DATA = ROOT / "raw", ROOT / "data"
DATA.mkdir(exist_ok=True)

PAGINA_SERIES = (
    "https://www.poderjudicial.es/cgpj/es/Temas/Estadistica-Judicial/Estudios-e-Informes/"
    "Informes-por-territorios-sobre-la-actividad-de-los-organos-judiciales/"
    "Informes-por-territorios-sobre-la-actividad-de-los-organos-judiciales"
)
BASE_XLSX = "https://www.poderjudicial.es/stfls/ESTADISTICA/FICHEROS/3002 Series estadisticas/"

ORDENES = {  # hoja en ficheros CGPJ -> nombre del orden jurisdiccional
    "Serie Total": "total",
    "Serie Civil": "civil",
    "Serie Penal": "penal",
    "Serie Contencioso": "contencioso",
    "Serie Contenciosa": "contencioso",
    "Serie Social": "social",
    "Serie total": "total",
    "Serie civil": "civil",
    "Serie penal": "penal",
    "Serie contencioso": "contencioso",
    "Serie social": "social",
}

stats = {"celdas_leidas": 0, "celdas_vacias": 0, "filas_descartadas": []}


def sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {name: list(wb[name].iter_rows(values_only=True)) for name in wb.sheetnames}
    wb.close()
    return out


def intro_text(rows):
    txt = []
    for row in rows:
        t = " ".join(str(v) for v in row if v is not None).strip()
        if t:
            txt.append(t)
    return txt


def parse_asuntos(rows):
    """Fichero 'Series Asuntos': bloques Ingresados/Resueltos/En trámite por territorio."""
    data = {}
    year_headers = {}
    current = None
    for row in rows:
        label = row[1]
        if label in ("Ingresados", "Resueltos", "En trámite", "En Trámite"):
            current = "en_tramite" if "trámite" in str(label).lower() else (
                "ingresados" if label == "Ingresados" else "resueltos")
            years = [int(v) for v in row[2:] if isinstance(v, (int, float))]
            year_headers[current] = years
            continue
        if current and isinstance(label, str) and label.strip() and not isinstance(row[2], str):
            years = year_headers[current]
            vals = row[2:2 + len(years)]
            entry = data.setdefault(label.strip(), {})
            for y, v in zip(years, vals):
                stats["celdas_leidas"] += 1
                if v is None:
                    stats["celdas_vacias"] += 1
                entry.setdefault(y, {})[current] = v
    return data


def parse_tasas(rows):
    """Ficheros de tasas: fila 8 con años, filas siguientes territorio -> valor/año."""
    data, years = {}, []
    for row in rows:
        if not years and isinstance(row[2], (int, float)):
            years = [int(v) for v in row[2:] if isinstance(v, (int, float))]
            continue
        name = row[1]
        if isinstance(name, str) and name.strip() and years and row[2] is not None \
                and not isinstance(row[2], str):
            vals = row[2:2 + len(years)]
            data[name.strip()] = {y: v for y, v in zip(years, vals) if v is not None}
            stats["celdas_leidas"] += len([v for v in vals if v is not None])
    return data



def add_per_capita(e, y, nombre, pop):
    """Pendientes por habitante.

    - Si INE publica población a 1 de enero del año (Tempus tabla 31304; llega hasta
      2022): en_tramite_por_habitante = en_tramite_final / población.
    - Si no (años 2023+): en_tramite_por_1000_hab_derivada = (en_tramite/ingresados)
      x tasa_litigiosidad, todo publicado por el CGPJ (equivale a la población que el
      propio CGPJ usa para su litigiosidad). DERIVADO, etiquetado como tal.
    """
    if not e.get("en_tramite_final"):
        return
    n_ine = nombre
    if n_ine and n_ine in pop and y in pop[n_ine]:
        e["en_tramite_por_habitante"] = round(
            e["en_tramite_final"] / pop[n_ine][y]["valor"], 4)
    elif e.get("tasa_litigiosidad") and e.get("ingresados"):
        e["en_tramite_por_1000_hab_derivada"] = round(
            e["en_tramite_final"] / e["ingresados"] * e["tasa_litigiosidad"], 2)


def main() -> int:



    # ---------- CGPJ: asuntos nacional (por orden y CCAA) ----------
    sh = sheets(RAW / "Series_Asuntos.xlsx")
    asuntos = {ORDENES[s]: parse_asuntos(rows) for s, rows in sh.items() if s in ORDENES}
    intro_asuntos = intro_text(sh["Introducción"])

    # ---------- CGPJ: tasas nacional (por orden y CCAA) ----------
    tasas = {}
    intros = {}
    for fname, key in [
        ("Series_Tasas_de_Congestion.xlsx", "tasa_congestion"),
        ("Series_Tasas_de_Resolucion.xlsx", "tasa_resolucion"),
        ("Series_Tasas_de_Pendencia.xlsx", "tasa_pendencia"),
        ("Series_Tasa_Litigiosidad.xlsx", "tasa_litigiosidad"),
    ]:
        sh = sheets(RAW / fname)
        tasas[key] = {ORDENES[s]: parse_tasas(rows) for s, rows in sh.items() if s in ORDENES}
        intros[key] = intro_text(sh["Introducción"])

    # ---------- CGPJ: provincias ----------
    sh_prov = sheets(RAW / "Series_Asuntos_por_Provincias.xlsx")
    asuntos_prov = parse_asuntos(sh_prov["Serie Total"])
    tasas_prov = {}
    for fname, key in [
        ("Series_Tasas_de_Congestion_por_Provincias.xlsx", "tasa_congestion"),
        ("Series_Tasas_de_Resolucion_por_Provincias.xlsx", "tasa_resolucion"),
        ("Series_Tasas_de_Pendencia_por_Provincias.xlsx", "tasa_pendencia"),
        ("Series_Tasa_Litigiosidad_por_Provincias.xlsx", "tasa_litigiosidad"),
    ]:
        sh = sheets(RAW / fname)
        hoja = next(n for n in sh if n.lower().startswith("serie"))
        tasas_prov[key] = parse_tasas(sh[hoja])

    # ---------- CGPJ: evolución de la pendencia (variación interanual publicada) ----------
    sh_evo = sheets(RAW / "Evolución_Asuntos_en_trámite_final_Período_Total_nacional.xlsx")
    evolucion = {}
    for row in sh_evo["Evolución Total"]:
        if isinstance(row[1], (int, float)) and row[2] is not None:
            evolucion[int(row[1])] = row[2]

    # ---------- INE: población ----------
    ine = json.loads((RAW / "ine_poblacion.json").read_text())
    pop = {}
    for territorio, puntos in ine.items():
        for p in puntos:
            pop.setdefault(territorio, {})[p["anyo"]] = {
                "valor": p["valor"], "tipo_dato": p["tipo_dato"]}

    # Mapeo de nombres INE -> nombres CGPJ (donde difieren)
    alias = {"Asturias, Principado de": "Asturias",
             "Comunitat Valenciana": "Comunitat Valenciana",
             "Madrid, Comunidad de": "Madrid",
             "Murcia, Región de": "Murcia",
             "Navarra, Comunidad Foral de": "Navarra",
             "Rioja, La": "La Rioja"}
    ine_names = {}
    for n in pop:
        ine_names[alias.get(n, n)] = n

    meta_cgpj = {
        "source": "CGPJ, Sección de Estadística Judicial. Series estadísticas",
        "url": BASE_XLSX, "pagina": PAGINA_SERIES, "unit": "asuntos",
        "period": "2001-2025, datos anuales",
        "nota": "En los datos de España no están incluidos los órganos centrales (Tribunal "
                "Supremo, Audiencia Nacional, juzgados centrales), según avisa el propio fichero.",
    }

    # ---------- data/nacional_anual.json ----------
    nac = {}
    for orden, d in asuntos.items():
        esp = d.get("España", {})
        for y, v in esp.items():
            nac.setdefault(str(y), {})["anyo"] = y
            for campo in ("ingresados", "resueltos", "en_tramite"):
                if v.get(campo) is not None:
                    nac[str(y)][f"{orden}_{campo}"] = v[campo]
    for key, d in tasas.items():
        esp = d.get("total", {}).get("España", {})
        for y, v in esp.items():
            nac[str(y)][key] = round(v, 4)
    out = {
        "meta": {**meta_cgpj,
                 "tasa_definiciones": "ver data/sources.json",
                 "unit_mixta": "asuntos para *_ingresados/*_resueltos/*_en_tramite; "
                               "tasas adimensionales salvo litigiosidad (asuntos/1.000 hab)"},
        "evolucion_pendencia_publicada": {
            "meta": {**meta_cgpj, "unit": "variación interanual de asuntos en trámite (fracción)",
                     "period": "1996-2025"},
            "datos": evolucion,
        },
        "nacional": nac,
    }
    (DATA / "nacional_anual.json").write_text(json.dumps(out, ensure_ascii=False, indent=1))

    # ---------- data/nacional_por_orden.json (2021-2025, últimos 5 años) ----------
    por_orden = {}
    for y in range(2021, 2026):
        anyo = {}
        for orden in ("civil", "penal", "contencioso", "social"):
            esp = asuntos[orden].get("España", {}).get(y, {})
            anyo[orden] = {
                "ingresados": esp.get("ingresados"),
                "resueltos": esp.get("resueltos"),
                "en_tramite_final": esp.get("en_tramite"),
                "tasa_congestion": tasas["tasa_congestion"][orden].get("España", {}).get(y),
                "tasa_resolucion": tasas["tasa_resolucion"][orden].get("España", {}).get(y),
                "tasa_pendencia": tasas["tasa_pendencia"][orden].get("España", {}).get(y),
            }
        por_orden[str(y)] = anyo
    (DATA / "nacional_por_orden.json").write_text(json.dumps(
        {"meta": meta_cgpj, "por_orden": por_orden}, ensure_ascii=False, indent=1))

    # ---------- data/ccaa.json (total jurisdicciones, 2001-2025) ----------
    ccaa_out = {}
    for ccaa in asuntos["total"]:
        if ccaa == "España":
            continue
        serie = {}
        for y, v in asuntos["total"][ccaa].items():
            e = {"ingresados": v.get("ingresados"), "resueltos": v.get("resueltos"),
                 "en_tramite_final": v.get("en_tramite")}
            for key in tasas:
                val = tasas[key]["total"].get(ccaa, {}).get(y)
                if val is not None:
                    e[key] = round(val, 4)
            add_per_capita(e, y, ine_names.get(ccaa), pop)
            serie[str(y)] = e
        ccaa_out[ccaa] = serie
    (DATA / "ccaa.json").write_text(json.dumps({
        "meta": {
            **meta_cgpj,
            "en_tramite_por_habitante": {
                "source": "DERIVADO: en_tramite_final (CGPJ) ÷ población a 1 de enero (INE, "
                          "tabla Tempus 31304, ambos sexos, total edad)",
                "unit": "asuntos en trámite por habitante",
                "nota": "Población INE con cifras provisionales o definitivas según año; "
                        "el CGPJ usa población INE para su tasa de litigiosidad.",
            },
        },
        "comunidades": ccaa_out}, ensure_ascii=False, indent=1))

    # ---------- data/provincias.json (total jurisdicciones, 2001-2025) ----------
    prov_out = {}
    for prov, serie in asuntos_prov.items():
        e_serie = {}
        for y, v in serie.items():
            e = {"ingresados": v.get("ingresados"), "resueltos": v.get("resueltos"),
                 "en_tramite_final": v.get("en_tramite")}
            for key in tasas_prov:
                val = tasas_prov[key].get(prov, {}).get(y)
                if val is not None:
                    e[key] = round(val, 4)
            add_per_capita(e, y, ine_names.get(prov), pop)
            e_serie[str(y)] = e
        prov_out[prov] = e_serie
    (DATA / "provincias.json").write_text(json.dumps({
        "meta": {
            **meta_cgpj,
            "nota": "Series por provincias solo para el total de jurisdicciones (los ficheros "
                    "por provincia no desglosan por orden en la hoja de serie histórica).",
            "en_tramite_por_habitante": {
                "source": "DERIVADO: en_tramite_final (CGPJ) ÷ población a 1 de enero (INE)",
                "unit": "asuntos en trámite por habitante"},
        },
        "provincias": prov_out}, ensure_ascii=False, indent=1))

    # ---------- data/sources.json ----------
    def find_def(texts, *keys):
        for t in texts:
            low = t.lower()
            if any(k in low for k in keys):
                return t
        return None

    defs = {
        "tasa_congestion": find_def(intros.get("tasa_congestion", []), "congestión", "congestion"),
        "tasa_resolucion": find_def(intros.get("tasa_resolucion", []), "resolución", "resolucion"),
        "tasa_pendencia": find_def(intros.get("tasa_pendencia", []), "pendencia"),
        "tasa_litigiosidad": find_def(intros.get("tasa_litigiosidad", []), "litigiosidad"),
    }
    (DATA / "sources.json").write_text(json.dumps({
        "fuente_principal": {
            "nombre": "CGPJ — Informes por territorios sobre la actividad de los órganos "
                      "judiciales → Series estadísticas (XLSX)",
            "url": PAGINA_SERIES,
            "formato": "XLSX, una hoja por orden jurisdiccional + hojas anuales 2001-2025",
            "granularidad": "Nacional + comunidad autónoma (por orden jurisdiccional); "
                            "provincias (total jurisdicciones)",
            "ficheros": {f: BASE_XLSX + f.replace("_", " ") for f in [
                "Series_Asuntos.xlsx", "Series_Asuntos_por_Provincias.xlsx",
                "Series_Tasas_de_Congestion.xlsx", "Series_Tasas_de_Congestion_por_Provincias.xlsx",
                "Series_Tasas_de_Resolucion.xlsx", "Series_Tasas_de_Resolucion_por_Provincias.xlsx",
                "Series_Tasas_de_Pendencia.xlsx", "Series_Tasas_de_Pendencia_por_Provincias.xlsx",
                "Series_Tasa_Litigiosidad.xlsx", "Series_Tasa_Litigiosidad_por_Provincias.xlsx"]},
        },
        "definiciones_oficiales_tasas": defs,
        "definiciones_aviso": [
            "Los ficheros 'Tasas de Congestion por Provincias' y 'Tasas de Pendencia por "
            "Provincias' llevan en su hoja Introducción la definición de la tasa de RESOLUCIÓN "
            "(error de copy-paste del CGPJ); sus valores sí corresponden a la tasa del título "
            "del fichero (verificado numéricamente en el informe).",
        ],
        "nota_espana": intro_asuntos,
        "fuente_poblacion": {
            "nombre": "INE — Cifras de población a 1 de enero (tabla Tempus 31304)",
            "url": "https://servicios.ine.es/wstempus/js/ES/DATOS_TABLA/31304?det=0&tip=AM",
            "uso": "solo para el derivado en_tramite_por_habitante; la tasa de litigiosidad "
                   "usa la que publica el propio CGPJ",
        },
        "duracion_media": {
            "publicada": "sí, pero NO en estas series XLSX",
            "punto_de_entrada": "https://www.poderjudicial.es/cgpj/es/Temas/Transparencia/"
                                "ch.Estimacion-de-los-tiempos-medios-de-duracion-de-los-"
                                "procedimientos-judiciales.formato1?anio=2024",
            "nota": "Herramienta interactiva por tipo de órgano, orden, procedimiento, año y "
                    "territorio (últimos 10 años, en meses). El CGPJ avisa de que es una "
                    "ESTIMACIÓN obtenida por modelo matemático a partir de ingresados, "
                    "resueltos y en trámite, no un dato observado. No hay descarga masiva; "
                    "no se agrega en este observatorio.",
        },
    }, ensure_ascii=False, indent=1))

    # ---------- resumen de conteo para el informe ----------
    resumen = {
        "celdas_leidas": stats["celdas_leidas"],
        "celdas_vacias": stats["celdas_vacias"],
        "comunidades": len(ccaa_out),
        "provincias": len(prov_out),
        "anyos": sorted(int(y) for y in nac),
    }
    (DATA / "_build_resumen.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=1))
    print(json.dumps(resumen, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
